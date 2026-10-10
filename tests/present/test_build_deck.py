import sys
from pathlib import Path
import pytest
from pptx import Presentation
from pptx.enum.text import PP_ALIGN

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "skills/present/references"))
import build_deck

def spec(slides, mode="live", **kw):
    return {"title": "Test deck", "mode": mode, "slides": slides, **kw}

def cnvpr(shape):
    return shape._element.xpath("./*[1]/p:cNvPr")[0]

def test_titles_and_notes(tmp_path):
    out = tmp_path / "d.pptx"
    build_deck.build(spec([
        {"layout": "title", "title": "Fieldnotes diary study", "notes": "Open with the problem."},
        {"layout": "claim_bullets", "title": "Technicians write notes after the job",
         "bullets": ["7 of 9 wrote in the van"], "notes": "Claim, then P4's quote."},
    ]), out)
    prs = Presentation(out)
    assert [s.shapes.title.text for s in prs.slides] == ["Fieldnotes diary study", "Technicians write notes after the job"]
    assert prs.slides[1].notes_slide.notes_text_frame.text == "Claim, then P4's quote."

def test_nested_bullets_rejected(tmp_path):
    with pytest.raises(ValueError, match="flat"):
        build_deck.build(spec([{"layout": "claim_bullets", "title": "A claim title here", "bullets": ["a", ["b"]]}]), tmp_path / "d.pptx")

def test_unknown_layout_rejected(tmp_path):
    with pytest.raises(ValueError, match="layout"):
        build_deck.build(spec([{"layout": "agenda", "title": "Agenda"}]), tmp_path / "d.pptx")

def test_missing_title_rejected(tmp_path):
    with pytest.raises(ValueError, match="title"):
        build_deck.build(spec([{"layout": "claim_bullets", "title": " ", "bullets": ["a"]}]), tmp_path / "d.pptx")

def test_visual_requires_alt(tmp_path):
    with pytest.raises(ValueError, match="alt"):
        build_deck.build(spec([{"layout": "claim_visual", "title": "A claim title here", "placeholder": "SCREENSHOT: step 2"}]), tmp_path / "d.pptx")

def test_placeholder_frame_carries_label_and_alt(tmp_path):
    out = tmp_path / "d.pptx"
    build_deck.build(spec([{"layout": "claim_visual", "title": "Option B removes the ambiguity",
        "placeholder": "SCREENSHOT: step 2, option B", "alt": "Visit type chosen before times appear"}]), out)
    shapes = list(Presentation(out).slides[0].shapes)
    frame = next(s for s in shapes if s.has_text_frame and "SCREENSHOT" in s.text_frame.text)
    assert frame.text_frame.text == "[SCREENSHOT: step 2, option B]"
    assert cnvpr(frame).get("descr") == "Visit type chosen before times appear"

def test_bar_chart_starts_at_zero(tmp_path):
    out = tmp_path / "d.pptx"
    build_deck.build(spec([{"layout": "claim_chart", "title": "Downgrade tickets dominate billing support",
        "chart": {"type": "bar", "categories": ["Downgrade", "Refund"], "series": [{"name": "Tickets", "values": [1240, 310]}], "highlight": 0},
        "alt": "Downgrade tickets 1,240 a month versus 310 refunds"}]), out)
    ch = next(s for s in Presentation(out).slides[0].shapes if s.has_chart)
    assert ch.chart.value_axis.minimum_scale == 0
    assert cnvpr(ch).get("descr").startswith("Downgrade tickets")

def test_read_mode_body_text(tmp_path):
    out = tmp_path / "d.pptx"
    build_deck.build(spec([{"layout": "claim_bullets", "title": "Self-serve downgrades cut support load",
        "body": "Customers must contact support today.", "bullets": ["1,240 tickets a month"]}], mode="read"), out)
    text = " ".join(s.text_frame.text for s in Presentation(out).slides[0].shapes if s.has_text_frame)
    assert "Customers must contact support today." in text

def test_template_layouts_used(tmp_path):
    tpl = tmp_path / "tpl.pptx"
    p = Presentation(); p.save(tpl)
    out = tmp_path / "d.pptx"
    build_deck.build(spec([{"layout": "title", "title": "Templated deck"}], template=str(tpl)), out)
    prs = Presentation(out)
    assert prs.slides[0].slide_layout.name == "Title Slide"
    assert prs.slides[0].shapes.title.text == "Templated deck"

def test_visual_and_bullets_do_not_overlap(tmp_path):
    out = tmp_path / "d.pptx"
    build_deck.build(spec([{"layout": "claim_visual", "title": "Option B removes the ambiguity",
        "placeholder": "SCREENSHOT: step 2", "alt": "Step two screen", "bullets": ["one", "two"]}]), out)
    shapes = list(Presentation(out).slides[0].shapes)
    frame = next(s for s in shapes if s.has_text_frame and "SCREENSHOT" in s.text_frame.text)
    bul = next(s for s in shapes if s.has_text_frame and s.text_frame.text.startswith("•"))
    assert frame.left + frame.width <= bul.left

def test_template_without_title_placeholder_rejected(tmp_path):
    # Layout removal is feasible via SlideLayouts.remove (python-pptx 1.0).
    p = Presentation()
    for lay in list(p.slide_layouts):
        if lay.name != "Blank":
            p.slide_layouts.remove(lay)
    tpl = tmp_path / "blank.pptx"; p.save(tpl)
    with pytest.raises(ValueError, match="title placeholder"):
        build_deck.build(spec([{"layout": "title", "title": "X deck"}], template=str(tpl)), tmp_path / "d.pptx")

def test_default_title_spans_ten_columns(tmp_path):
    out = tmp_path / "d.pptx"
    build_deck.build(spec([{"layout": "claim_bullets", "title": "Wide title", "bullets": ["a"]}]), out)
    t = Presentation(out).slides[0].shapes.title
    assert (t.left, t.width) == build_deck.span(1, 10)

def test_chart_spec_validated(tmp_path):
    bad = {"type": "bar", "categories": ["a", "b"], "series": [{"name": "s", "values": [1]}]}
    with pytest.raises(ValueError, match="slide 1"):
        build_deck.build(spec([{"layout": "claim_chart", "title": "A claim title here", "chart": bad, "alt": "x"}]), tmp_path / "d.pptx")
    with pytest.raises(ValueError, match="chart"):
        build_deck.build(spec([{"layout": "claim_chart", "title": "A claim title here", "chart": {"type": "pie"}, "alt": "x"}]), tmp_path / "d.pptx")

def test_chart_labeled_directly_without_legend(tmp_path):
    out = tmp_path / "d.pptx"
    build_deck.build(spec([{"layout": "claim_chart", "title": "Redesign cut errors in both regions",
        "chart": {"type": "bar", "categories": ["EU", "US"],
                  "series": [{"name": "After", "values": [12, 9]}, {"name": "Before", "values": [30, 28]}], "highlight": 0},
        "alt": "Errors fell in both regions after the redesign"}], mode="read"), out)
    ch = next(s for s in Presentation(out).slides[0].shapes if s.has_chart).chart
    assert len(list(ch.series)) == 2
    assert ch.has_legend is False
    plot = ch.plots[0]
    assert plot.has_data_labels
    dl = plot.data_labels
    assert dl.show_value
    assert dl.number_format == "General"
    assert dl.font.size.pt == build_deck.SIZES["read"]["small"]

def _body_sizes(out, text):
    shp = next(s for s in Presentation(out).slides[0].shapes if s.has_text_frame and s.text_frame.text == text)
    return {r.font.size.pt for p in shp.text_frame.paragraphs for r in p.runs}

def test_projected_read_deck_uses_projection_sizes(tmp_path):
    out = tmp_path / "d.pptx"
    build_deck.build(spec([{"layout": "claim_bullets", "title": "Self-serve downgrades cut support load",
        "body": "Customers must contact support today."}], mode="read", projected=True), out)
    assert _body_sizes(out, "Customers must contact support today.") == {18}

def test_plain_read_deck_keeps_read_sizes(tmp_path):
    out = tmp_path / "d.pptx"
    build_deck.build(spec([{"layout": "claim_bullets", "title": "Self-serve downgrades cut support load",
        "body": "Customers must contact support today."}], mode="read"), out)
    assert _body_sizes(out, "Customers must contact support today.") == {14}

def test_projected_must_be_bool(tmp_path):
    with pytest.raises(ValueError, match="projected"):
        build_deck.build(spec([{"layout": "title", "title": "X deck"}], mode="read", projected="yes"), tmp_path / "d.pptx")

def _chart_slide(chart):
    return spec([{"layout": "claim_chart", "title": "A claim title here", "chart": chart, "alt": "x"}])

def test_chart_values_must_be_numbers(tmp_path):
    for bad in (["12", 9], [True, 9], [None, 9]):
        c = {"type": "bar", "categories": ["a", "b"], "series": [{"name": "s", "values": bad}]}
        with pytest.raises(ValueError, match="slide 1"):
            build_deck.build(_chart_slide(c), tmp_path / "d.pptx")

def test_chart_highlight_must_index_a_series(tmp_path):
    for hi in (1, -1, "0", True, 0.0):
        c = {"type": "bar", "categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2.5]}], "highlight": hi}
        with pytest.raises(ValueError, match="highlight"):
            build_deck.build(_chart_slide(c), tmp_path / "d.pptx")


# --- The grid: same margins and 12-column structure on every slide ---

def _png(path, w=400, h=300):
    from PIL import Image
    Image.new("RGB", (w, h), (200, 200, 220)).save(path)
    return str(path)

def _mixed_slides(img):
    chart = {"type": "bar", "categories": ["A", "B"], "series": [{"name": "n", "values": [4, 2]}]}
    body = "What the presenter would have said, in two or three sentences."
    return [
        {"layout": "title", "title": "The deck's governing message", "subtitle": "Review group, Friday", "body": body},
        {"layout": "section", "title": "The second half turns on cost", "subtitle": "Why it matters"},
        {"layout": "claim_visual", "title": "Option B removes the ambiguity", "placeholder": "SCREENSHOT: step 2", "alt": "x", "body": body},
        {"layout": "claim_visual", "title": "Option B with its reasons", "placeholder": "SCREENSHOT: step 2", "alt": "x", "bullets": ["one", "two"], "body": body},
        {"layout": "claim_visual", "title": "A wide screenshot fits its span", "image": _png(img, 2400, 400), "alt": "x"},
        {"layout": "claim_visual", "title": "A tall screenshot fits the body", "image": _png(img.with_name("tall.png"), 300, 1200), "alt": "x", "bullets": ["a"]},
        {"layout": "claim_chart", "title": "Downgrades dominate tickets", "chart": chart, "alt": "x", "body": body},
        {"layout": "claim_chart", "title": "Downgrades with context", "chart": chart, "alt": "x", "bullets": ["a", "b"], "subtitle": "n = 412", "body": body},
        {"layout": "claim_bullets", "title": "Technicians write notes later", "bullets": ["7 of 9 in the van"], "subtitle": "Diary study", "body": body},
        {"layout": "quote", "title": "Notes are written from memory", "quote": "I write it all up in the van after.", "attribution": "P4, technician", "subtitle": "Diary study", "body": body},
        {"layout": "ask", "title": "Approve direction B so engineering can start", "subtitle": "Review group, by Friday", "body": body},
    ]

def _grid_decks(tmp_path):
    img = tmp_path / "wide.png"
    slides = _mixed_slides(img)
    for kw in ({"mode": "live"}, {"mode": "read"}, {"mode": "read", "projected": True}):
        live = kw["mode"] == "live"
        out = tmp_path / f"{'-'.join(map(str, kw.values()))}.pptx"
        build_deck.build(spec([{k: v for k, v in s.items() if not (live and k == "body")} for s in slides], **kw), out)
        yield Presentation(out)

def test_grid_span_math():
    px = build_deck._px
    assert build_deck.span(1, 12) == (px(80), px(1120))
    assert build_deck.span(8, 5) == (px(80 + 7 * 96), px(5 * 64 + 4 * 32))
    assert (build_deck.MARGIN_X, build_deck.MARGIN_Y) == (px(80), px(64))
    assert (build_deck.TITLE_TOP, build_deck.BODY_TOP, build_deck.BODY_BOTTOM) == (px(64), px(224), px(656))

def test_every_shape_inside_the_content_box(tmp_path):
    px = build_deck._px
    for prs in _grid_decks(tmp_path):
        assert (prs.slide_width, prs.slide_height) == (build_deck.SLIDE_W, build_deck.SLIDE_H)
        for n, slide in enumerate(prs.slides, 1):
            for shp in slide.shapes:
                assert shp.left >= px(80) and shp.top >= px(64), (n, shp.name)
                assert shp.left + shp.width <= px(1200), (n, shp.name)
                assert shp.top + shp.height <= px(656), (n, shp.name)

def test_titles_sit_in_the_title_band(tmp_path):
    for prs in _grid_decks(tmp_path):
        for slide, s in zip(prs.slides, _mixed_slides(tmp_path / "wide.png")):
            t = slide.shapes.title
            assert t.left == build_deck.MARGIN_X
            hero = s["layout"] in ("title", "section")
            assert t.top == (build_deck.BODY_TOP if hero else build_deck.TITLE_TOP), s["layout"]
            assert t.width == build_deck.span(1, 10)[1]
            assert all(p.alignment == PP_ALIGN.LEFT for p in t.text_frame.paragraphs)

def _named(slide, pred):
    return next(s for s in slide.shapes if pred(s))

def test_visual_and_bullets_split_on_column_boundaries(tmp_path):
    for prs in _grid_decks(tmp_path):
        for idx in (3, 7):  # placeholder + bullets, chart + bullets
            slide = prs.slides[idx]
            vis = _named(slide, lambda s: s.has_chart or (s.has_text_frame and s.text_frame.text.startswith("[")))
            bul = _named(slide, lambda s: s.has_text_frame and s.text_frame.text.startswith("•"))
            assert (vis.left, vis.width) == build_deck.span(1, 7)
            assert (bul.left, bul.width) == build_deck.span(8, 5)
            assert vis.left + vis.width + build_deck.GUTTER == bul.left

def test_lone_visual_spans_all_twelve_columns(tmp_path):
    for prs in _grid_decks(tmp_path):
        for idx in (2, 6):
            vis = _named(prs.slides[idx], lambda s: s.has_chart or (s.has_text_frame and s.text_frame.text.startswith("[")))
            assert (vis.left, vis.width) == build_deck.span(1, 12)
            assert vis.top + vis.height == build_deck.BODY_BOTTOM

def test_claim_bullets_and_quote_keep_readable_spans(tmp_path):
    for prs in _grid_decks(tmp_path):
        bul = _named(prs.slides[8], lambda s: s.has_text_frame and s.text_frame.text.startswith("•"))
        assert bul.width == build_deck.span(1, 8)[1]
        q = _named(prs.slides[9], lambda s: s.has_text_frame and s.text_frame.text.startswith("“"))
        assert (q.left, q.width) == build_deck.span(1, 9)

def test_body_content_on_the_eight_unit_rhythm(tmp_path):
    for prs in _grid_decks(tmp_path):
        for slide in prs.slides:
            for shp in slide.shapes:
                if shp.top >= build_deck.BODY_TOP:
                    assert (shp.top - build_deck.BODY_TOP) % build_deck._px(8) == 0, shp.name

def test_read_body_text_tops_the_body_region(tmp_path):
    prs = list(_grid_decks(tmp_path))[1]
    body = _named(prs.slides[2], lambda s: s.has_text_frame and s.text_frame.text.startswith("What the presenter"))
    assert body.top == build_deck.BODY_TOP
    assert (body.left, body.width) == build_deck.span(1, 8)

def test_template_body_scales_with_the_slide(tmp_path):
    tpl = tmp_path / "tpl.pptx"
    Presentation().save(tpl)  # default python-pptx template: 10 x 7.5 in, 4:3
    out = tmp_path / "d.pptx"
    build_deck.build(spec([{"layout": "claim_bullets", "title": "Templated claim", "bullets": ["a"]}], template=str(tpl)), out)
    prs = Presentation(out)
    bul = _named(prs.slides[0], lambda s: s.has_text_frame and s.text_frame.text.startswith("•"))
    sx = prs.slide_width / build_deck.SLIDE_W
    left, width = build_deck.span(1, 8)
    assert abs(bul.left - left * sx) <= 1 and abs(bul.width - width * sx) <= 1
    assert bul.left + bul.width <= prs.slide_width - build_deck.MARGIN_X * sx + 1


def test_chart_has_no_automatic_title(tmp_path):
    out = tmp_path / "d.pptx"
    build_deck.build(spec([{"layout": "claim_chart", "title": "A third of sessions abandon at step 2",
        "chart": {"type": "bar", "categories": ["Step 1", "Step 2"], "series": [{"name": "Abandon %", "values": [8, 34]}]},
        "alt": "Step 2 abandons at 34%"}]), out)
    ch = next(s for s in Presentation(out).slides[0].shapes if s.has_chart)
    deleted = ch.chart._chartSpace.xpath('c:chart/c:autoTitleDeleted')
    assert deleted and deleted[0].get('val') in ('1', 'true')


def test_notes_over_three_sentences_rejected(tmp_path):
    notes = "One claim. Two reasons hold it. Three is the transition. Four is too many."
    with pytest.raises(ValueError, match="three sentences"):
        build_deck.build(spec([{"layout": "claim_bullets", "title": "A claim title goes right here",
            "bullets": ["a"], "notes": notes}]), tmp_path / "d.pptx")


def test_bulleted_notes_rejected(tmp_path):
    with pytest.raises(ValueError, match="plain prose"):
        build_deck.build(spec([{"layout": "claim_bullets", "title": "A claim title goes right here",
            "bullets": ["a"], "notes": "- point one\n- point two"}]), tmp_path / "d.pptx")


def test_three_sentence_notes_accepted(tmp_path):
    out = tmp_path / "d.pptx"
    notes = "Step 2 loses a third of sessions. Testers couldn't tell visit types apart. That's what option B fixes."
    build_deck.build(spec([{"layout": "claim_bullets", "title": "A claim title goes right here",
        "bullets": ["a"], "notes": notes}]), out)
    assert Presentation(out).slides[0].notes_slide.notes_text_frame.text == notes
