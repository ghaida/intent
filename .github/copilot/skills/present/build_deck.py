#!/usr/bin/env python3
"""Build a .pptx from a /present deck spec (JSON). Shared path for PowerPoint and
Google Slides (import the .pptx into Slides). Usage: build_deck.py deck.json out.pptx

Requires python-pptx (pip install python-pptx). Enforces the /present rules a
machine can check: every slide titled, flat bullets only, alt text on
every visual, bar charts from zero, speaker notes in the notes field, three sentences at most."""
import json, re, sys
from pptx import Presentation
from pptx.util import Emu, Inches, Pt
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER

LAYOUTS = {"title", "claim_visual", "claim_chart", "claim_bullets", "quote", "section", "ask"}
INK, MUTED, RULE, SURFACE = RGBColor(0x1A, 0x1A, 0x2E), RGBColor(0x5B, 0x5B, 0x72), RGBColor(0xB8, 0xB8, 0xC8), RGBColor(0xF1, 0xF1, 0xF6)
ACCENT = RGBColor(0x43, 0x38, 0xCA)
SIZES = {"live": {"title": 36, "body": 24, "small": 18}, "read": {"title": 28, "body": 14, "small": 12}}
PROJECTED_READ = {"title": 32, "body": 18, "small": 14}  # read layout, projection-safe sizes: nothing on a projected slide drops below 18pt body

# The grid. Every slide shares one frame: a 1280 x 720 stage (13.333 x 7.5 in at 96 units
# per inch), 80/64 margins, 12 columns of 64 with 32 gutters, a 128-tall title band and a
# body region below it. Values are stage units; _px converts to EMU exactly
# (914400 / 96 = 9525, so this equals Inches(n / 96) without float truncation).
def _px(n):
    return Emu(round(n * 9525))

SLIDE_W, SLIDE_H = _px(1280), _px(720)
MARGIN_X, MARGIN_Y = _px(80), _px(64)
COL, GUTTER = _px(64), _px(32)
TITLE_TOP, TITLE_H = _px(64), _px(128)
BODY_TOP, BODY_BOTTOM = _px(224), _px(656)
HERO_H, SUB_H, BODY_H, QUOTE_H, ATTR_H, SRC_H, STEP = 160, 48, 128, 192, 48, 32, 16  # stage units, all on the 8-unit rhythm

def span(first_col, n_cols, sx=1.0):
    """(left, width) in EMU of n_cols columns starting at first_col (1-based).
    sx scales the grid for a template whose slide isn't 13.333 in wide."""
    left = 80 + (first_col - 1) * 96
    width = n_cols * 64 + (n_cols - 1) * 32
    return _px(left * sx), _px(width * sx)

def _validate(spec):
    if spec.get("mode", "live") not in SIZES:
        raise ValueError("mode must be 'live' or 'read'")
    if not isinstance(spec.get("projected", False), bool):
        raise ValueError("projected must be true or false")
    for i, s in enumerate(spec.get("slides", []), 1):
        if s.get("layout") not in LAYOUTS:
            raise ValueError(f"slide {i}: unknown layout {s.get('layout')!r}; use one of {sorted(LAYOUTS)}")
        if not str(s.get("title", "")).strip():
            raise ValueError(f"slide {i}: every slide needs a title")
        for key in ("title", "subtitle", "body", "quote", "attribution", "source", "notes"):
            if "\u2014" in str(s.get(key, "")):
                raise ValueError(f"slide {i}: no em-dash in {key}; use a period, a colon or a comma")
        if any("\u2014" in str(x) for x in s.get("bullets", []) if isinstance(x, str)):
            raise ValueError(f"slide {i}: no em-dash in bullets; use a period, a colon or a comma")
        notes = str(s.get("notes", "")).strip()
        if notes:
            if re.search(r"^\s*([-*•]|\d+[.)])\s", notes, re.M):
                raise ValueError(f"slide {i}: notes are plain prose, no bullets or numbered lines")
            if len(re.findall(r"[^.!?]+[.!?]+(?=\s|$)", notes)) > 3:
                raise ValueError(f"slide {i}: notes hold three sentences at most; the presenter reads them at a glance")
        b = s.get("bullets", [])
        if not isinstance(b, list) or any(not isinstance(x, str) for x in b):
            raise ValueError(f"slide {i}: bullets must be a flat list of strings, no nesting")
        if (s.get("image") or s.get("placeholder") or s.get("chart")) and not str(s.get("alt", "")).strip():
            raise ValueError(f"slide {i}: alt text is required for every image, placeholder or chart")
        c = s.get("chart")
        if c is not None:
            if not isinstance(c, dict) or c.get("type") not in ("bar", "line"):
                raise ValueError(f"slide {i}: chart type must be 'bar' or 'line'")
            cats, series = c.get("categories"), c.get("series")
            if not isinstance(cats, list) or not cats:
                raise ValueError(f"slide {i}: chart needs a non-empty categories list")
            if not isinstance(series, list) or not series:
                raise ValueError(f"slide {i}: chart needs a non-empty series list")
            for ser in series:
                if not isinstance(ser, dict) or not str(ser.get("name", "")).strip() \
                        or not isinstance(ser.get("values"), list) or len(ser["values"]) != len(cats):
                    raise ValueError(f"slide {i}: each chart series needs a name and one value per category")
                if any(isinstance(v, bool) or not isinstance(v, (int, float)) for v in ser["values"]):
                    raise ValueError(f"slide {i}: chart values must be numbers, as stated in the source")
            hi = c.get("highlight", 0)
            if isinstance(hi, bool) or not isinstance(hi, int) or hi not in range(len(series)):
                raise ValueError(f"slide {i}: chart highlight must be the index of one series (0 to {len(series) - 1})")

def _set_alt(shape, alt):
    shape._element.xpath("./*[1]/p:cNvPr")[0].set("descr", alt)

def _flush(tf):
    # Zero insets, top-anchored: the text edge is the grid edge, not 0.1 in inside it.
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP

def _text(slide, text, left, top, width, height, size, color=INK, bold=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame; tf.word_wrap = True; _flush(tf)
    tf.text = text
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
    return tb

def _bullets(slide, items, left, top, width, height, size):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame; tf.word_wrap = True; _flush(tf)
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.level = f"• {item}", 0
        for r in p.runs:
            r.font.size, r.font.color.rgb = Pt(size), INK
    return tb

def _chart(slide, c, left, top, width, height, alt, label_size):
    data = CategoryChartData(); data.categories = c["categories"]
    for s in c["series"]:
        data.add_series(s["name"], s["values"])
    kind = XL_CHART_TYPE.COLUMN_CLUSTERED if c["type"] == "bar" else XL_CHART_TYPE.LINE
    gf = slide.shapes.add_chart(kind, left, top, width, height, data)
    ch = gf.chart
    if c["type"] == "bar":
        ch.value_axis.minimum_scale = 0
    ch.has_title = False   # the slide's claim title is the chart's title
    ch.has_legend = False  # label directly: values on the marks, series identity from the accent + the claim title
    plot = ch.plots[0]
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.show_value, dl.number_format, dl.number_format_is_linked = True, "General", False
    dl.font.size, dl.font.color.rgb = Pt(label_size), INK
    ch.value_axis.has_major_gridlines = False
    hi = c.get("highlight", 0)
    for i, series in enumerate(ch.series):
        color = ACCENT if i == hi else RULE
        if c["type"] == "bar":
            series.format.fill.solid(); series.format.fill.fore_color.rgb = color
        else:
            series.format.line.color.rgb = color
    _set_alt(gf, alt)
    return gf

def _placeholder(slide, label, alt, left, top, width, height):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid(); shp.fill.fore_color.rgb = SURFACE
    shp.line.color.rgb = RULE
    shp.text_frame.text = f"[{label}]"
    for r in shp.text_frame.paragraphs[0].runs:
        r.font.size, r.font.color.rgb = Pt(14), MUTED
    _set_alt(shp, alt)
    return shp

def _new_slide(prs, templated, layout):
    if templated:
        names = {l.name: l for l in prs.slide_layouts}
        if layout in ("title", "section", "ask"):
            lay = names.get("Title Slide") or names.get("Section Header") or prs.slide_layouts[0]
        else:
            lay = names.get("Title Only") or prs.slide_layouts[min(5, len(prs.slide_layouts) - 1)]
        has_title = lambda l: any(ph.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE) for ph in l.placeholders)
        if not has_title(lay):
            lay = next((l for l in prs.slide_layouts if has_title(l)), None)
            if lay is None:
                raise ValueError("template has no layout with a title placeholder")
        return prs.slides.add_slide(lay)
    return prs.slides.add_slide(prs.slide_layouts[5])  # default template "Title Only"

def _title(slide, s, size):
    t = slide.shapes.title
    t.text = s["title"]
    for p in t.text_frame.paragraphs:
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb = Pt(size), True, INK
    return t

def build(spec, out_path):
    _validate(spec)
    mode = spec.get("mode", "live")
    sz = PROJECTED_READ if mode == "read" and spec.get("projected") else SIZES[mode]
    templated = bool(spec.get("template"))
    prs = Presentation(spec["template"]) if templated else Presentation()
    if not templated:
        prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    # A template that isn't 13.333 x 7.5 in gets the same grid scaled to its slide: x by the
    # width ratio, y by the height ratio, so margins, columns and bands keep their share of
    # the slide. Its own title placeholders keep the template's geometry.
    sx, sy = prs.slide_width / SLIDE_W, prs.slide_height / SLIDE_H
    X = lambda first, n: span(first, n, sx)
    Y = lambda v: _px(v * sy)
    for s in spec["slides"]:
        lay = s["layout"]
        hero = lay in ("title", "section")
        slide = _new_slide(prs, templated, lay)
        t = _title(slide, s, sz["title"] + (8 if hero else 0))
        y = 224  # stage units; every offset below is a multiple of 8
        if not templated:
            _flush(t.text_frame)
            for p in t.text_frame.paragraphs:
                p.alignment = PP_ALIGN.LEFT  # the default master centers titles; the grid starts them at the margin
            t.left, t.width = span(1, 10)
            t.top, t.height = (BODY_TOP, _px(HERO_H)) if hero else (TITLE_TOP, TITLE_H)
            if hero:
                y += HERO_H
        else:  # never stack body content over the template's title, wherever it sits
            bottom = (t.top + t.height) / prs.slide_height * 720 + STEP
            y = max(y, -(-int(bottom) // 8) * 8)

        def place(h):  # claim the next block of the body region on the 8-unit rhythm
            nonlocal y
            top = y; y += h + STEP
            return Y(top), Y(h)

        if s.get("subtitle"):
            left, w = X(1, 10 if hero else 8)
            top, h = place(SUB_H)
            _text(slide, s["subtitle"], left, top, w, h, sz["small"], MUTED)
        if mode == "read" and s.get("body"):
            top, h = place(BODY_H)
            _text(slide, s["body"], X(1, 8)[0], top, X(1, 8)[1], h, sz["body"], INK)
        src = SRC_H + STEP if s.get("source") else 0  # the source line keeps the last band of the body region
        room = max(656 - src - y, 48)  # stage units from the cursor to the bottom of the body region
        top = Y(y)
        if lay == "quote" and s.get("quote"):
            qh = min(QUOTE_H, room - (ATTR_H + STEP if s.get("attribution") else 0))
            left, w = X(1, 9)
            _text(slide, f"“{s['quote']}”", left, top, w, Y(qh), sz["body"] + 4, INK)
            if s.get("attribution"):
                _text(slide, s["attribution"], left, Y(y + qh + STEP), w, Y(ATTR_H), sz["small"], MUTED)
        visual = s.get("chart") or s.get("image") or s.get("placeholder")
        both = bool(visual and s.get("bullets"))
        vleft, vw = X(1, 7) if both else X(1, 12)
        bleft, bw = X(8, 5) if both else X(1, 8)
        vh = Y(room)
        if s.get("chart"):
            _chart(slide, s["chart"], vleft, top, vw, vh, s["alt"], sz["small"])
        elif s.get("image"):
            pic = slide.shapes.add_picture(s["image"], vleft, top)
            k = min(vw / pic.width, vh / pic.height)  # fit the span and the body region, aspect kept
            pic.width, pic.height = int(pic.width * k), int(pic.height * k)
            _set_alt(pic, s["alt"])
        elif s.get("placeholder"):
            _placeholder(slide, s["placeholder"], s["alt"], vleft, top, vw, vh)
        if s.get("bullets"):
            _bullets(slide, s["bullets"], bleft, top, bw, vh, sz["body"])
        if s.get("source"):
            left, w = X(1, 12)
            _text(slide, s["source"], left, Y(656 - SRC_H), w, Y(SRC_H), sz["small"], MUTED)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
    prs.save(out_path)

if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("usage: build_deck.py deck.json out.pptx")
    build(json.load(open(sys.argv[1], encoding="utf-8")), sys.argv[2])
    print(f"wrote {sys.argv[2]}")
