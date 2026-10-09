#!/usr/bin/env python3
"""Build a .pptx from a /present deck spec (JSON). Shared path for PowerPoint and
Google Slides (import the .pptx into Slides). Usage: build_deck.py deck.json out.pptx

Requires python-pptx (pip install python-pptx). Enforces the /present rules that
can be enforced mechanically: every slide titled, flat bullets only, alt text on
every visual, bar charts from zero, speaker notes in the notes field."""
import json, sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER

LAYOUTS = {"title", "claim_visual", "claim_chart", "claim_bullets", "quote", "section", "ask"}
INK, MUTED, RULE, SURFACE = RGBColor(0x1A, 0x1A, 0x2E), RGBColor(0x5B, 0x5B, 0x72), RGBColor(0xB8, 0xB8, 0xC8), RGBColor(0xF1, 0xF1, 0xF6)
ACCENT = RGBColor(0x43, 0x38, 0xCA)
SIZES = {"live": {"title": 36, "body": 24, "small": 18}, "read": {"title": 28, "body": 14, "small": 12}}
W, H = Inches(13.333), Inches(7.5)
M = Inches(0.75)

def _validate(spec):
    if spec.get("mode", "live") not in SIZES:
        raise ValueError("mode must be 'live' or 'read'")
    for i, s in enumerate(spec.get("slides", []), 1):
        if s.get("layout") not in LAYOUTS:
            raise ValueError(f"slide {i}: unknown layout {s.get('layout')!r}; use one of {sorted(LAYOUTS)}")
        if not str(s.get("title", "")).strip():
            raise ValueError(f"slide {i}: every slide needs a title")
        b = s.get("bullets", [])
        if not isinstance(b, list) or any(not isinstance(x, str) for x in b):
            raise ValueError(f"slide {i}: bullets must be a flat list of strings — no nesting")
        if (s.get("image") or s.get("placeholder") or s.get("chart")) and not str(s.get("alt", "")).strip():
            raise ValueError(f"slide {i}: alt text is required for every image, placeholder or chart")
        c = s.get("chart")
        if c is not None:
            if not isinstance(c, dict) or c.get("type") not in ("bar", "line"):
                raise ValueError(f"slide {i}: chart type must be 'bar' or 'line'")
            cats, series = c.get("categories"), c.get("series")
            if not isinstance(cats, list) or not cats:
                raise ValueError(f"slide {i}: chart needs a non-empty categories list")
            if not isinstance(series, list) or not series:
                raise ValueError(f"slide {i}: chart needs a non-empty series list")
            for ser in series:
                if not isinstance(ser, dict) or not str(ser.get("name", "")).strip() \
                        or not isinstance(ser.get("values"), list) or len(ser["values"]) != len(cats):
                    raise ValueError(f"slide {i}: each chart series needs a name and one value per category")

def _set_alt(shape, alt):
    shape._element.xpath("./*[1]/p:cNvPr")[0].set("descr", alt)

def _text(slide, text, left, top, width, height, size, color=INK, bold=False):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame; tf.word_wrap = True
    tf.text = text
    for p in tf.paragraphs:
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
    return tb

def _bullets(slide, items, left, top, width, height, size):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.level = f"• {item}", 0
        for r in p.runs:
            r.font.size, r.font.color.rgb = Pt(size), INK
    return tb

def _chart(slide, c, left, top, width, height, alt):
    data = CategoryChartData(); data.categories = c["categories"]
    for s in c["series"]:
        data.add_series(s["name"], s["values"])
    kind = XL_CHART_TYPE.COLUMN_CLUSTERED if c["type"] == "bar" else XL_CHART_TYPE.LINE
    gf = slide.shapes.add_chart(kind, left, top, width, height, data)
    ch = gf.chart
    if c["type"] == "bar":
        ch.value_axis.minimum_scale = 0
    ch.has_legend = len(c["series"]) > 1
    if ch.has_legend:
        ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    ch.value_axis.has_major_gridlines = False
    hi = c.get("highlight", 0)
    for i, series in enumerate(ch.series):
        color = ACCENT if i == hi else RULE
        if c["type"] == "bar":
            series.format.fill.solid(); series.format.fill.fore_color.rgb = color
        else:
            series.format.line.color.rgb = color
    _set_alt(gf, alt)
    return gf

def _placeholder(slide, label, alt, left, top, width, height):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shp.fill.solid(); shp.fill.fore_color.rgb = SURFACE
    shp.line.color.rgb = RULE
    shp.text_frame.text = f"[{label}]"
    for r in shp.text_frame.paragraphs[0].runs:
        r.font.size, r.font.color.rgb = Pt(14), MUTED
    _set_alt(shp, alt)
    return shp

def _new_slide(prs, templated, layout):
    if templated:
        names = {l.name: l for l in prs.slide_layouts}
        if layout in ("title", "section", "ask"):
            lay = names.get("Title Slide") or names.get("Section Header") or prs.slide_layouts[0]
        else:
            lay = names.get("Title Only") or prs.slide_layouts[min(5, len(prs.slide_layouts) - 1)]
        has_title = lambda l: any(ph.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE) for ph in l.placeholders)
        if not has_title(lay):
            lay = next((l for l in prs.slide_layouts if has_title(l)), None)
            if lay is None:
                raise ValueError("template has no layout with a title placeholder")
        return prs.slides.add_slide(lay)
    return prs.slides.add_slide(prs.slide_layouts[5])  # default template "Title Only"

def _title(slide, s, size):
    t = slide.shapes.title
    t.text = s["title"]
    for p in t.text_frame.paragraphs:
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb = Pt(size), True, INK
    return t

def build(spec, out_path):
    _validate(spec)
    mode = spec.get("mode", "live"); sz = SIZES[mode]
    templated = bool(spec.get("template"))
    prs = Presentation(spec["template"]) if templated else Presentation()
    if not templated:
        prs.slide_width, prs.slide_height = W, H
    cw = prs.slide_width - 2 * M
    for s in spec["slides"]:
        lay = s["layout"]
        slide = _new_slide(prs, templated, lay)
        t = _title(slide, s, sz["title"] + (8 if lay in ("title", "section") else 0))
        if not templated:
            t.left, t.top, t.width, t.height = M, Inches(0.5), cw, Inches(1.2)
        y = Inches(1.9)
        if s.get("subtitle"):
            _text(slide, s["subtitle"], M, y, cw, Inches(0.8), sz["small"], MUTED); y += Inches(0.9)
        if mode == "read" and s.get("body"):
            _text(slide, s["body"], M, y, cw, Inches(1.2), sz["body"], INK); y += Inches(1.3)
        room = prs.slide_height - y - Inches(0.6)
        if lay == "quote" and s.get("quote"):
            _text(slide, f"“{s['quote']}”", M, y, cw, Inches(2), sz["body"] + 4, INK)
            if s.get("attribution"):
                _text(slide, f"— {s['attribution']}", M, y + Inches(2.1), cw, Inches(0.6), sz["small"], MUTED)
        visual = s.get("chart") or s.get("image") or s.get("placeholder")
        both = bool(visual and s.get("bullets"))
        vw = int(cw * 0.6) if both else cw
        bleft = M + vw + Inches(0.3) if both else M
        bw = cw - vw - Inches(0.3) if both else cw
        if s.get("chart"):
            _chart(slide, s["chart"], M, y, vw, room, s["alt"])
        elif s.get("image"):
            pic = slide.shapes.add_picture(s["image"], M, y, height=room)
            if pic.width > vw:  # keep aspect ratio, fit the allotted width
                pic.height = int(pic.height * vw / pic.width); pic.width = vw
            _set_alt(pic, s["alt"])
        elif s.get("placeholder"):
            _placeholder(slide, s["placeholder"], s["alt"], M, y, vw, room)
        if s.get("bullets"):
            _bullets(slide, s["bullets"], bleft, y, bw, room, sz["body"])
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = s["notes"]
    prs.save(out_path)

if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("usage: build_deck.py deck.json out.pptx")
    build(json.load(open(sys.argv[1], encoding="utf-8")), sys.argv[2])
    print(f"wrote {sys.argv[2]}")
